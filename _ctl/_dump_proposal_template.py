from docx import Document
import os, json
base=r'C:\Users\thang\Downloads\CAPSTONE_TEMPLATE_2026\PROPOSAL_TEMPLATE'
files=[
'2026 -SE_Capstone_Proposal_Template.docx',
'Guided_SE&NE_Capstone_Proposal_Template(xem hưởng dẫn ở đây).docx'
]
for fn in files:
    path=os.path.join(base,fn)
    print('\n### FILE:',fn)
    doc=Document(path)
    print('sections',len(doc.sections),'paragraphs',len(doc.paragraphs),'tables',len(doc.tables))
    for i,p in enumerate(doc.paragraphs):
        text=p.text.replace('\t',' <TAB> ').strip()
        if text:
            print(f'P{i:03d} [{p.style.name}] {text}')
    for ti,t in enumerate(doc.tables):
        print(f'\nTABLE {ti} rows={len(t.rows)} cols={len(t.columns)}')
        for ri,row in enumerate(t.rows):
            vals=[]
            for ci,cell in enumerate(row.cells):
                txt=' | '.join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                vals.append(txt)
            print(f'R{ri:02d}: '+ ' || '.join(vals))
